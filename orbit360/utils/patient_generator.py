"""
patient_generator.py — Orbit360

Synthesizes test-patient data from a multi-sheet Excel template.
Refactored from the original interactive `tools/generate_patients.py` so the
GUI (Genesis panel) and the CLI can both drive the same logic.

Output layout (one .xlsx per source sheet so each batch flows straight into
the existing single-sheet pool model used by excel_data_manager):

    orbit_data/test_data/generated/<RUN_ID>/
        <template_stem>_<release>_<sheet>.xlsx   (one per template sheet)
        csv/<sheet>_<RUN_ID>.csv
        patient_generation_<RUN_ID>.log
"""
from __future__ import annotations

import csv
import logging
import random
import re
import string
from dataclasses import dataclass, field
from datetime import datetime, timedelta
from pathlib import Path
from typing import Callable, Literal

from openpyxl import Workbook, load_workbook

from orbit360.utils.paths import GENERATED_PATIENTS_DIR, PATIENT_TEMPLATES_DIR

_LOG = logging.getLogger(__name__)

SystemType = Literal["EXPANSE", "MAGIC"]

ProgressCallback = Callable[[str, int, int], None]
"""Called with (sheet_name, completed_in_sheet, total_in_sheet)."""


@dataclass
class GenerateResult:
    output_dir: Path
    pool_files: list[Path] = field(default_factory=list)   # one .xlsx per source sheet
    csv_files: list[Path]  = field(default_factory=list)
    log_file: Path | None  = None
    total_patients: int    = 0
    sheet_count: int       = 0
    duration_secs: float   = 0.0
    system: SystemType     = "MAGIC"


# ── Discovery / template inspection ─────────────────────────────────────────

def list_templates() -> list[Path]:
    """All .xlsx templates the user has dropped under PATIENT_TEMPLATES_DIR."""
    if not PATIENT_TEMPLATES_DIR.is_dir():
        return []
    return sorted(p for p in PATIENT_TEMPLATES_DIR.glob("*.xlsx") if p.is_file())


def detect_system(template_path: Path) -> SystemType:
    """Match the original CLI rule: 'MTX' in upper(name) -> EXPANSE."""
    return "EXPANSE" if "MTX" in template_path.name.upper() else "MAGIC"


@dataclass
class TemplateSummary:
    path: Path
    system: SystemType
    sheet_names: list[str]
    column_count: int   # column count of the first sheet


def inspect_template(template_path: Path) -> TemplateSummary:
    """Lightweight read of headers only, used for the GUI status line."""
    wb = load_workbook(template_path, read_only=True, data_only=True)
    try:
        sheet_names = list(wb.sheetnames)
        col_count = 0
        if sheet_names:
            ws = wb[sheet_names[0]]
            for row in ws.iter_rows(min_row=1, max_row=1, values_only=True):
                col_count = sum(1 for c in row if c is not None) or len(row)
                break
        return TemplateSummary(
            path=template_path,
            system=detect_system(template_path),
            sheet_names=sheet_names,
            column_count=col_count,
        )
    finally:
        wb.close()


# ── Field synthesis (lifted verbatim from the original CLI) ─────────────────

_FIRST_NAMES = ["James", "Michael", "Robert", "John", "David"]
_LAST_NAMES  = ["Smith", "Johnson", "Williams", "Brown", "Jones"]
_STREET_NAMES = ["Main", "Oak", "Pine", "Maple", "Cedar", "Elm", "Lake", "Hill", "View", "Park"]
_STREET_TYPES = ["St", "Ave", "Rd", "Ln", "Dr"]


def _suffix(index: int) -> str:
    letters = string.ascii_uppercase
    result = ""
    while True:
        result = letters[index % 26] + result
        index = index // 26 - 1
        if index < 0:
            break
    return result


def _build_name(system: SystemType, patient_type: str, release: str, index: int) -> str:
    suf = _suffix(index)
    if system == "EXPANSE":
        return f"ABS3M,{patient_type} {release}.{suf}"
    return f"{patient_type},{release}.{suf}"


def _random_address() -> str:
    return f"{random.randint(100, 9999)} {random.choice(_STREET_NAMES)} {random.choice(_STREET_TYPES)}"


def _random_dob() -> str:
    start = datetime(1940, 1, 1)
    end   = datetime(2010, 1, 1)
    return (start + timedelta(days=random.randint(0, (end - start).days))).strftime("%m%d%y")


def _random_ssn() -> str:
    return str(random.randint(100_000_000, 999_999_999))


def _random_phone() -> str:
    return str(random.randint(2_000_000_000, 9_999_999_999))


def _generate_value(
    field: str,
    template_value,
    patient_name: str,
    person: dict,
    system: SystemType,
):
    field_lower = (field or "").lower()
    first = person["first"]
    last  = person["last"]

    if field_lower == "patientname":
        return patient_name

    if system == "EXPANSE":
        if "subscriber" in field_lower or "guarantor" in field_lower:
            return patient_name
        if "firstname" in field_lower:
            return first
        if "lastname" in field_lower:
            return last

    if system == "MAGIC":
        if "nextofkin" in field_lower:
            return f"{last},{first}"

    if "address" in field_lower:
        return _random_address()
    if "dob" in field_lower:
        return _random_dob()
    if "ssn" in field_lower or "social" in field_lower:
        return _random_ssn()
    if "phone" in field_lower:
        return _random_phone()

    if any(token in field_lower for token in ("mrn", "acct", "account", "visit")):
        return ""

    return template_value


# ── Public entry point ──────────────────────────────────────────────────────

_RELEASE_RE = re.compile(r"[^A-Z0-9._-]+")


def _slug_release(release: str) -> str:
    return _RELEASE_RE.sub("_", release.strip().upper()).strip("_") or "REL"


def _slug_sheet(sheet_name: str) -> str:
    return re.sub(r"[^A-Za-z0-9._-]+", "_", sheet_name).strip("_") or "SHEET"


def make_run_id() -> str:
    """Match the orbit_logger run-id format so generated batch dirs sort
    naturally next to runs/."""
    return datetime.now().strftime("%Y-%m-%d_%H-%M-%S")


def generate_patients(
    template_path: Path,
    release: str,
    count_per_sheet: int,
    output_root: Path | None = None,
    progress: ProgressCallback | None = None,
) -> GenerateResult:
    """Generate a fresh batch of test patients from *template_path*.

    Each source sheet becomes its own .xlsx pool file under
    output_root/<RUN_ID>/, plus a CSV mirror and a log of the run.
    """
    template_path = Path(template_path)
    if not template_path.is_file():
        raise FileNotFoundError(f"Template not found: {template_path}")
    if count_per_sheet <= 0:
        raise ValueError("count_per_sheet must be > 0")
    release = release.strip()
    if not release:
        raise ValueError("release must be non-empty")

    output_root = Path(output_root) if output_root else GENERATED_PATIENTS_DIR
    run_id = make_run_id()
    batch_dir = output_root / run_id
    csv_dir   = batch_dir / "csv"
    batch_dir.mkdir(parents=True, exist_ok=True)
    csv_dir.mkdir(parents=True, exist_ok=True)

    log_file = batch_dir / f"patient_generation_{run_id}.log"

    def _log(msg: str, level: str = "INFO") -> None:
        ts = datetime.now().strftime("%H:%M:%S")
        line = f"[{ts}] | {level:<7} | {msg}"
        with log_file.open("a", encoding="utf-8") as fh:
            fh.write(line + "\n")

    system  = detect_system(template_path)
    release_slug = _slug_release(release)

    started = datetime.now()
    _log(f"START template={template_path.name} release={release} system={system} count={count_per_sheet}")

    wb_in = load_workbook(template_path, data_only=True)
    try:
        result = GenerateResult(output_dir=batch_dir, log_file=log_file, system=system)
        result.sheet_count = len(wb_in.sheetnames)

        for sheet_name in wb_in.sheetnames:
            ws_in   = wb_in[sheet_name]
            rows    = list(ws_in.iter_rows(values_only=True))
            if not rows:
                _log(f"SKIP sheet={sheet_name} (empty)", level="WARNING")
                continue

            headers      = list(rows[0])
            template_row = list(rows[1]) if len(rows) > 1 else [""] * len(headers)

            # One workbook per source sheet so the existing single-sheet pool
            # model in excel_data_manager picks each up as its own pool.
            wb_out = Workbook()
            ws_out = wb_out.active
            ws_out.title = sheet_name[:31] or "Sheet"  # Excel sheet name limit
            ws_out.append(headers)

            csv_rows: list[list] = []
            patient_type = sheet_name.upper()

            for i in range(count_per_sheet):
                name  = _build_name(system, patient_type, release, i)
                first = random.choice(_FIRST_NAMES)
                last  = random.choice(_LAST_NAMES)
                person = {"first": first, "last": last}

                new_row = [
                    _generate_value(headers[col], template_row[col] if col < len(template_row) else "",
                                    name, person, system)
                    for col in range(len(headers))
                ]
                ws_out.append(new_row)
                csv_rows.append(new_row)

                if progress and (i + 1) % 25 == 0:
                    progress(sheet_name, i + 1, count_per_sheet)

            # Final progress tick for the sheet
            if progress:
                progress(sheet_name, count_per_sheet, count_per_sheet)

            pool_name = f"{template_path.stem}_{release_slug}_{_slug_sheet(sheet_name)}.xlsx"
            pool_path = batch_dir / pool_name
            wb_out.save(pool_path)
            wb_out.close()
            result.pool_files.append(pool_path)

            csv_path = csv_dir / f"{_slug_sheet(sheet_name)}_{run_id}.csv"
            with csv_path.open("w", newline="", encoding="utf-8") as fh:
                writer = csv.writer(fh)
                writer.writerow(headers)
                writer.writerows(csv_rows)
            result.csv_files.append(csv_path)

            result.total_patients += len(csv_rows)
            _log(f"DONE   sheet={sheet_name} patients={len(csv_rows)} pool={pool_name}")
    finally:
        wb_in.close()

    result.duration_secs = (datetime.now() - started).total_seconds()
    _log(f"COMPLETE total={result.total_patients} sheets={result.sheet_count} secs={result.duration_secs:.2f}")
    return result
